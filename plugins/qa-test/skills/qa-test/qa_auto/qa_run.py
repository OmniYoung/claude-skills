"""
QA 자동화 러너 - spec.json을 읽어 실제 사이트에서 테스트를 실행하고 리포트를 생성한다.

사용법:
  py <skill>/qa_auto/qa_run.py --init .            작업 폴더(qa_auto/) 생성
  py <skill>/qa_auto/qa_run.py qa_auto/specs/xxx.json
  py <skill>/qa_auto/qa_run.py ... --dry-run       쓰기 스텝 전부 건너뜀 (셀렉터 검증용)
  py <skill>/qa_auto/qa_run.py ... --headed        브라우저 띄워서 눈으로 확인
  py <skill>/qa_auto/qa_run.py ... --case TC-003   특정 케이스만
  py <skill>/qa_auto/qa_run.py --login URL [URL2 ...] --ws qa_auto
                                                   로그인 창. 로그인하면 저장 후 저절로 닫힘

산출물 위치:
  이 스크립트는 스킬 폴더에 아무것도 쓰지 않는다. 스펙 파일 위치로 작업 폴더를 정한다.
  <작업폴더>/specs/xxx.json 이면 <작업폴더> 아래 output/ 과 fixtures/ 를 쓴다.
  덕분에 팀원 각자가 자기 프로젝트에서 같은 스킬을 공유해 쓸 수 있다.

안전장치:
  - 쓰기 스텝은 케이스에 approved:true 가 있어야만 실행된다 (없으면 SKIP)
  - blocklist 매칭 시 승인 여부와 무관하게 케이스 중단
  - 쓰기 케이스는 cleanup 스텝이 없으면 실행 거부
"""

import argparse
import json
import pathlib
import re
import sys
import time
from datetime import datetime
from pathlib import Path

from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

# 한국어 윈도우 콘솔은 기본이 cp949 라, 케이스 제목이나 실패 사유에 들어간
# "—" "·" "※" 같은 문자에서 출력 단계가 죽는다. 리포트(UTF-8)는 멀쩡한데
# 콘솔 print 하나 때문에 실행 전체가 중단되므로 진입점에서 맞춰둔다.
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

sys.path.insert(0, str(Path(__file__).parent))
from report import build_report, _dev_items

# SKILL_DIR: 스킬에 번들된 스크립트·기본 픽스처가 있는 곳. 여기에는 아무것도 쓰지 않는다.
SKILL_DIR = Path(__file__).parent


def resolve_workspace(spec_path):
    """스펙 위치로 작업 폴더를 정한다. <ws>/specs/x.json 이면 <ws>, 아니면 스펙이 있는 폴더."""
    return spec_path.parent.parent if spec_path.parent.name == "specs" else spec_path.parent


def find_fixture(workspace, name):
    """작업 폴더의 fixtures 를 먼저 보고, 없으면 스킬 번들 기본 픽스처로 넘어간다."""
    for base in (workspace / "fixtures", SKILL_DIR / "fixtures"):
        f = (base / name).resolve()
        if f.exists():
            return f
    raise FileNotFoundError(
        "업로드 파일 없음: {} (찾은 곳: {}, {})".format(
            name, workspace / "fixtures", SKILL_DIR / "fixtures"))

# -- 쓰기 판정 --------------------------------------------------
WRITE_ACTIONS = {"fill", "type", "select", "check", "uncheck", "upload"}
WRITE_HINTS = ["저장", "확인", "등록", "추가", "삭제", "수정", "전송", "발송", "적용", "완료"]
READ_ACTIONS = {
    "goto", "wait", "screenshot", "hover",
    "expect_visible", "expect_hidden", "expect_text",
    "expect_count", "expect_attr", "expect_dialog", "expect_no_dialog",
    "expect_js", "download", "expect_file",
}

# expect_js 는 읽기 전용이어야 한다. 아래 패턴이 보이면 실행을 거부해
# 임의 JS로 쓰기 승인 게이트를 우회하는 길을 막는다.
JS_MUTATION_PATTERNS = [
    ".click(", ".submit(", ".focus(", "fetch(", "XMLHttpRequest", "navigator.send",
    ".value=", ".value =", ".innerHTML=", ".innerHTML =", ".remove(", ".setAttribute(",
    "localStorage", "sessionStorage", "document.cookie", "location=", "location.href",
    "dispatchEvent", "eval(",
]
ALL_ACTIONS = WRITE_ACTIONS | READ_ACTIONS | {"click", "dblclick", "press"}

# press 로 누르는 키 중 상태를 바꾸지 않는 것들. 나머지(Enter 등)는 쓰기로 본다.
READ_KEYS = {"escape", "tab", "arrowup", "arrowdown", "arrowleft", "arrowright",
             "pageup", "pagedown", "home", "end", "shift+tab"}

DEFAULT_VIEWPORT = {"width": 1600, "height": 900}


def is_write_step(step):
    """스텝이 서버에 쓰기를 하는지 판정. spec_schema.md의 '쓰기 판정 규칙' 구현."""
    action = step.get("action")
    if "write" in step:                      # 규칙 3,4 - 명시값이 최우선
        return bool(step["write"])
    if action in WRITE_ACTIONS:              # 규칙 1
        return True
    if action in ("click", "dblclick"):       # 규칙 2 - 텍스트 힌트로 자동 승격
        target = "{} {}".format(step.get("selector", ""), step.get("label", ""))
        return any(hint in target for hint in WRITE_HINTS)
    if action == "press":                     # Enter 는 저장, Esc·Tab 등은 이동일 뿐
        return step.get("key", "").lower() not in READ_KEYS
    return False


def hits_blocklist(step, blocklist):
    """blocklist 패턴에 걸리는 스텝인지. 걸리면 승인 여부 무관하게 차단."""
    target = "{} {} {}".format(
        step.get("selector", ""), step.get("label", ""), step.get("value", "")
    )
    for pattern in blocklist:
        needle = pattern.split("=", 1)[1] if pattern.startswith("text=") else pattern
        if needle and needle in target:
            return pattern
    return None


# -- 쿠키 ------------------------------------------------------

class SessionExpired(Exception):
    pass


def load_cookies(spec_path, rel, workspace):
    """{"domain": "...", "cookies": [{"name":..., "value":...}]} 포맷.

    경로는 스펙 파일 기준으로 먼저 찾고, 없으면 작업 폴더 기준으로 찾는다.
    덕분에 스펙에 "cookies.json" 한 줄만 써도 <작업폴더>/cookies.json 을 집어온다.
    """
    for base in (spec_path.parent, workspace):
        path = (base / rel).resolve()
        if path.exists():
            break
    else:
        path = (spec_path.parent / rel).resolve()
    if not path.exists():
        raise FileNotFoundError(
            "쿠키 파일 없음: {} (찾은 곳: {}, {})".format(rel, spec_path.parent, workspace))

    text = path.read_text(encoding="utf-8-sig").strip()
    if text.startswith("{"):
        return _cookies_from_json(text)
    return _cookies_from_devtools(text, path)


def _cookies_from_json(text):
    """{"domain": "...", "cookies": [{"name":..., "value":...}]} 포맷."""
    data = json.loads(text)
    return [
        {"name": c["name"], "value": c["value"], "domain": data["domain"], "path": "/"}
        for c in data["cookies"] if c.get("value")
    ]


def _cookies_from_devtools(text, path):
    """크롬 개발자도구 > Application > Cookies 의 표를 그대로 붙여넣은 탭 구분 텍스트.

    컬럼 순서: Name / Value / Domain / Path / Expires / Size / ... 뒤는 무시한다.
    값이 비었거나 컬럼이 모자란 줄은 건너뛴다. 헤더 줄도 자동으로 건너뛴다.
    """
    out, skipped = [], 0
    for line in text.splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue               # 빈 줄과 안내 주석
        col = line.split("\t")
        if len(col) < 4:
            skipped += 1
            continue
        name, value, domain, cpath = (c.strip() for c in col[:4])
        if name.lower() == "name" and value.lower() == "value":
            continue                      # 헤더 줄
        if not name or not value or not domain:
            skipped += 1
            continue
        out.append({"name": name, "value": value,
                    "domain": domain, "path": cpath or "/"})
    if not out:
        raise ValueError(
            "쿠키를 한 건도 못 읽음: {}\n"
            "크롬 F12 > Application > Cookies 에서 표 전체를 선택해 복사한 뒤 "
            "그대로 붙여넣어야 한다 (탭 구분).".format(path))
    if skipped:
        print("[*] 쿠키 {}건 로드 (형식이 안 맞는 {}줄 건너뜀)".format(len(out), skipped))
    else:
        print("[*] 쿠키 {}건 로드".format(len(out)))
    return out


# -- 다운로드 파일 읽기 -----------------------------------------

def read_table(path):
    """다운로드한 표 파일을 [[셀,...], ...] 로 읽는다.

    .xlsx 는 openpyxl, csv/tsv 는 구분자 분리, .xls 는 서버가 HTML 표를 그 확장자로
    내보내는 경우가 흔해서 태그를 벗겨 읽는다.
    """
    raw = path.read_bytes()
    suffix = path.suffix.lower()

    if raw[:2] == b"PK" and suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise StepFailure(
                "xlsx 를 읽으려면 openpyxl 이 필요하다: pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [["" if c is None else str(c).strip() for c in r]
                for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows

    text = raw.decode("utf-8-sig", errors="replace")
    if "<table" in text.lower():                      # HTML 표를 .xls 로 내보낸 경우
        rows = []
        for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", text, re.S | re.I):
            cells = [re.sub(r"<[^>]+>", "", c).replace("&nbsp;", " ").strip()
                     for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.S | re.I)]
            if cells:
                rows.append(cells)
        return rows

    sep = "\t" if "\t" in text.split("\n")[0] else ","
    return [[c.strip().strip('"') for c in line.split(sep)]
            for line in text.splitlines() if line.strip()]


# -- 스텝 실행 --------------------------------------------------

class StepFailure(Exception):
    pass


class Runner:
    def __init__(self, page, spec, spec_path, shot_dir, workspace):
        self.page = page
        self.spec = spec
        self.spec_path = spec_path
        self.shot_dir = shot_dir
        self.workspace = workspace
        # {QA_ROOT} = 번들된 스크립트 폴더의 file:// URI. 자체 점검용 로컬 픽스처를 가리킬 때 쓴다.
        self.base_url = spec["meta"]["base_url"].replace(
            "{QA_ROOT}", SKILL_DIR.as_uri()).rstrip("/")
        self.qa_prefix = spec["meta"].get("qa_prefix", "")
        self.dialogs = []          # 직전 액션이 띄운 알럿 버퍼
        self.warnings = []
        self.shots = []
        self.files = {}   # download 로 받은 파일
        page.on("dialog", self._on_dialog)

    def _on_dialog(self, dialog):
        self.dialogs.append(dialog.message)
        dialog.accept()

    def _check_session(self):
        """로그인이 풀렸는지 본다.

        세션이 만료되면 모든 케이스가 엉뚱한 이유로 줄줄이 실패해서 원인 파악이 오래 걸린다.
        첫 페이지에서 바로 잡아 중단시키는 편이 낫다.
        """
        chk = self.spec["meta"].get("login_check")
        if not chk:
            return
        url_hit = chk.get("url_contains") and chk["url_contains"] in self.page.url
        body_hit = False
        if chk.get("body_contains"):
            try:
                body_hit = chk["body_contains"] in (self.page.inner_text("body") or "")
            except Exception:
                body_hit = False
        if url_hit or body_hit:
            raise SessionExpired(
                "로그인 세션이 만료된 것으로 보인다 (현재 URL: {}). "
                "인증 파일을 갱신한 뒤 다시 실행할 것.".format(self.page.url))

    def _warn_prefix(self, value, case_id, idx):
        """테스트로 만든 데이터인지 눈으로 구분되게 QA 프리픽스를 권한다.

        검색어 입력처럼 프리픽스가 없어야 정상인 경우도 있어 차단은 하지 않고 경고만 남긴다.
        """
        if (self.qa_prefix and len(value) >= 2 and self.qa_prefix not in value
                and not value.replace(".", "").isdigit()):
            self.warnings.append(
                "{} 스텝{}: 입력값 '{}'에 QA 프리픽스({})가 없음 - 운영 데이터 오염 주의".format(
                    case_id, idx, value, self.qa_prefix))

    def _url(self, url):
        return url if url.startswith("http") else self.base_url + "/" + url.lstrip("/")

    def _shoot(self, name, step=None):
        """스크린샷. step 에 selector 가 있으면 그 지점만 잘라 찍는다.

        전체 화면만 찍으면 수백 행짜리 표에서 문제 행을 찾을 수가 없어 결함 증빙이 안 된다.
        clip 으로 잘라내므로 마우스를 움직이지 않는다 - 호버로 띄운 툴팁이 살아 있다.
        """
        safe = re.sub(r'[\\/:*?"<>|]', "_", name)
        path = self.shot_dir / (safe + ".png")
        clip = self._clip_for(step) if step else None
        if clip:
            self.page.screenshot(path=str(path), clip=clip)
        else:
            self.page.screenshot(path=str(path), full_page=True)
        self.shots.append(path.name)
        return path.name

    def _clip_for(self, step):
        """selector + 주변 맥락(context/include_header)을 감싸는 사각형을 구한다."""
        sel = step.get("selector")
        if not sel:
            return None
        target = self.page.locator(sel).first
        try:
            target.scroll_into_view_if_needed(timeout=5000)
            self.page.wait_for_timeout(120)
            b = target.bounding_box()
        except Exception:
            return None
        if not b:
            return None
        boxes = [b]

        # 위아래 정상 행을 같이 담아야 "이 행만 다르다"가 보인다
        ctx = step.get("context")
        if ctx:
            sib = ctx.get("selector", sel)
            try:
                pos, total = self.page.evaluate(
                    """([sel, sib]) => {
                        const all = Array.from(document.querySelectorAll(sib));
                        const t = document.querySelector(sel);
                        return [all.indexOf(t), all.length];
                    }""", [sel, sib])
            except Exception:
                pos, total = -1, 0
            if pos >= 0:
                lo = max(0, pos - int(ctx.get("above", 0)))
                hi = min(total - 1, pos + int(ctx.get("below", 0)))
                for i in (lo, hi):
                    try:
                        nb = self.page.locator(sib).nth(i).bounding_box()
                        if nb:
                            boxes.append(nb)
                    except Exception:
                        pass

        if step.get("include_header"):
            try:
                hb = self.page.locator(step["include_header"]).first.bounding_box()
                if hb:
                    boxes.append(hb)
            except Exception:
                pass

        pad = int(step.get("padding", 8))
        x0 = min(v["x"] for v in boxes) - pad
        y0 = min(v["y"] for v in boxes) - pad
        x1 = max(v["x"] + v["width"] for v in boxes) + pad
        y1 = max(v["y"] + v["height"] for v in boxes) + pad
        vw = self.page.viewport_size or {"width": 1600, "height": 900}
        return {"x": max(0, x0), "y": max(0, y0),
                "width": min(x1 - x0, vw["width"]), "height": y1 - y0}

    def run(self, step, case_id, idx):
        action = step.get("action")
        if action not in ALL_ACTIONS:
            raise StepFailure("정의되지 않은 action: {}".format(action))

        sel = step.get("selector")
        # expect_dialog 는 직전 버퍼를 봐야 하므로 여기서 비우면 안 된다
        if action not in ("expect_dialog", "expect_no_dialog"):
            self.dialogs = []

        if action == "goto":
            self.page.goto(self._url(step["url"]), wait_until="domcontentloaded")
            self.page.wait_for_timeout(600)
            self._check_session()

        elif action == "wait":
            if sel:
                self.page.wait_for_selector(sel, timeout=step.get("timeout", 10000))
            else:
                self.page.wait_for_timeout(step.get("ms", 1000))

        elif action == "download":
            # 엑셀 내보내기는 관리자 화면 단골 요건인데 "자동화 불가"로 빠지기 쉽다.
            # 받아서 파싱하면 행수·컬럼·순서까지 화면과 대조된다.
            dl_dir = self.shot_dir.parent / "downloads"
            dl_dir.mkdir(parents=True, exist_ok=True)
            with self.page.expect_download(timeout=step.get("timeout", 120000)) as info:
                self.page.click(sel, timeout=step.get("click_timeout", 10000))
            dl = info.value
            name = step.get("save_as") or dl.suggested_filename or "download"
            if "." not in name:
                name += pathlib.PurePath(dl.suggested_filename or "x.bin").suffix
            target = dl_dir / re.sub(r'[\\/:*?"<>|]', "_", name)
            dl.save_as(str(target))
            self.files[step.get("save_as") or name] = target
            return None

        elif action == "expect_file":
            key = step["file"]
            path = self.files.get(key)
            if path is None or not path.exists():
                raise StepFailure("받아둔 파일이 없음: {} (download 스텝이 먼저 와야 한다)".format(key))
            rows = read_table(path)
            if not rows:
                raise StepFailure("파일에서 행을 못 읽음: {}".format(path.name))

            if step.get("headers"):
                got = [c.replace(" ", "") for c in rows[0]]
                want = [c.replace(" ", "") for c in step["headers"]]
                if got[:len(want)] != want:
                    raise StepFailure("헤더 불일치 - 기대:{} / 실제:{}".format(want, got[:len(want)]))

            body = len(rows) - (1 if step.get("headers") else 0)
            if "min_rows" in step and body < step["min_rows"]:
                raise StepFailure("행 수 미달 - 기대:{}행 이상 / 실제:{}행".format(step["min_rows"], body))
            if "row_count_js" in step:
                # 운영 데이터는 계속 바뀌므로 기대값을 화면에서 읽어 비교한다
                expected = self.page.evaluate(step["row_count_js"])
                if body != expected:
                    raise StepFailure(
                        "행 수 불일치 - 화면 {}행 / 파일 {}행".format(expected, body))
            return None

        elif action == "screenshot":
            return self._shoot("{}_{:02d}_{}".format(case_id, idx, step.get("name", "shot")), step)

        elif action == "click":
            self.page.click(sel, timeout=step.get("timeout", 10000))
            self.page.wait_for_timeout(500)

        elif action == "dblclick":
            # 인라인 편집은 보통 더블클릭으로 input 이 생긴다. fill 은 그 input 을
            # 먼저 찾아야 해서 진입 자체가 안 되므로 별도 액션이 필요하다.
            self.page.dblclick(sel, timeout=step.get("timeout", 10000))
            self.page.wait_for_timeout(400)

        elif action == "press":
            key = step["key"]
            if sel:
                self.page.press(sel, key, timeout=step.get("timeout", 10000))
            else:
                self.page.keyboard.press(key)
            self.page.wait_for_timeout(400)

        elif action == "type":
            # fill 과 달리 기존 값을 지우지 않고 키 입력을 흉내낸다.
            # 글자수 제한처럼 입력 중 동작하는 검증에는 이쪽이 맞다.
            value = step.get("text", "")
            self._warn_prefix(value, case_id, idx)
            if step.get("clear"):
                self.page.fill(sel, "")
            self.page.type(sel, value, delay=step.get("delay", 20))

        elif action == "hover":
            # locator.hover()는 셀 안 인라인 span에서 hit-test를 오판해 타임아웃이 난다.
            # 요소 중심 좌표로 직접 마우스를 옮기는 방식이 안정적이다.
            loc = self.page.locator(sel).first
            # scroll_into_view_if_needed 는 "보이기만 하면" 멈춘다. 요소가 뷰포트
            # 하단에 걸쳐 있으면 아래로 열리는 툴팁이 렌더 영역 밖이라 거짓 실패가 난다.
            loc.scroll_into_view_if_needed(timeout=step.get("timeout", 10000))
            loc.evaluate("e => e.scrollIntoView({block:'center', inline:'center'})")
            self.page.wait_for_timeout(150)
            box = loc.bounding_box()
            if not box:
                raise StepFailure("호버 대상의 위치를 잡을 수 없음: {}".format(sel))
            self.page.mouse.move(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
            self.page.wait_for_timeout(step.get("ms", 800))

        elif action == "fill":
            value = step.get("value", "")
            self._warn_prefix(value, case_id, idx)
            self.page.fill(sel, value)

        elif action == "select":
            self.page.select_option(sel, step["value"])

        elif action == "check":
            self.page.check(sel)

        elif action == "uncheck":
            self.page.uncheck(sel)

        elif action == "upload":
            try:
                f = find_fixture(self.workspace, step["file"])
            except FileNotFoundError as e:
                raise StepFailure(str(e))
            self.page.set_input_files(sel, str(f))
            self.page.wait_for_timeout(800)

        elif action == "expect_visible":
            if not self.page.is_visible(sel):
                raise StepFailure("보여야 할 요소가 안 보임: {}".format(sel))

        elif action == "expect_hidden":
            if self.page.is_visible(sel):
                raise StepFailure("숨겨져야 할 요소가 보임: {}".format(sel))

        elif action == "expect_text":
            actual = (self.page.text_content(sel, timeout=step.get("timeout", 10000)) or "").strip()
            if "contains" in step and step["contains"] not in actual:
                raise StepFailure("문구 불일치 - 기대:'{}' / 실제:'{}'".format(
                    step["contains"], actual[:120]))
            if "not_contains" in step and step["not_contains"] in actual:
                raise StepFailure("있으면 안 되는 문구가 있음 - '{}' / 실제:'{}'".format(
                    step["not_contains"], actual[:150]))

        elif action == "expect_count":
            n = self.page.locator(sel).count()
            if "count" in step and n != step["count"]:
                raise StepFailure("개수 불일치 - 기대:{} / 실제:{}".format(step["count"], n))
            if "min" in step and n < step["min"]:
                raise StepFailure("최소 개수 미달 - 기대:{}개 이상 / 실제:{}".format(step["min"], n))

        elif action == "expect_attr":
            val = self.page.get_attribute(sel, step["attr"]) or ""
            if "equals" in step and val != step["equals"]:
                raise StepFailure("{} 불일치 - 기대:'{}' / 실제:'{}'".format(
                    step["attr"], step["equals"], val))
            if "contains" in step and step["contains"] not in val:
                raise StepFailure("{}에 '{}' 없음 - 실제:'{}'".format(
                    step["attr"], step["contains"], val))

        elif action == "expect_dialog":
            if not self.dialogs:
                raise StepFailure("알럿이 떠야 하는데 안 뜸 - 기대 문구:'{}'".format(step["contains"]))
            joined = " | ".join(self.dialogs)
            if step["contains"] not in joined:
                raise StepFailure("알럿 문구 불일치 - 기대:'{}' / 실제:'{}'".format(
                    step["contains"], joined))

        elif action == "expect_no_dialog":
            if self.dialogs:
                raise StepFailure("알럿이 뜨면 안 되는데 뜸: {}".format(" | ".join(self.dialogs)))

        elif action == "expect_js":
            # CSS로 표현 안 되는 판정(노출 순서·계산값 비교)에만 쓴다. 읽기 전용 강제.
            js = step["js"]
            flat = js.replace(" ", "")
            for bad in JS_MUTATION_PATTERNS:
                if bad.replace(" ", "") in flat:
                    raise StepFailure(
                        "expect_js 는 읽기 전용만 허용 - 금지 패턴 '{}' 발견".format(bad))
            value = self.page.evaluate(js)

            # {ok, detail} 을 돌려주면 ok 로 판정하고 detail 을 사유에 그대로 싣는다.
            # 기대값이 데이터에 따라 변해 equals 에 못 박는 경우(건수 비교 등)를 위한 것 —
            # 불리언만 쓰면 "기대:True / 실제:False" 로 남아 개발팀에 그대로 못 보낸다.
            if isinstance(value, dict) and "ok" in value:
                if not value["ok"]:
                    detail = value.get("detail", "")
                    raise StepFailure("{}{}".format(
                        step.get("desc", "expect_js 불일치"),
                        " - {}".format(detail) if detail else ""))
                return None

            expected = step.get("equals", True)
            if value != expected:
                raise StepFailure("{} - 기대:{} / 실제:{}".format(
                    step.get("desc", "expect_js 불일치"), expected, value))

        return None


# -- 케이스 실행 ------------------------------------------------

INTERACTION_ACTIONS = {"click", "fill", "select", "check", "uncheck", "upload", "hover"}


def diagnose_chain(runner, case, result):
    """chains 케이스가 실패했을 때, 이어서 해서 깨진 건지 원래 깨진 건지 가린다.

    실패한 스텝을 **직전 동작 하나만 앞세워** 새 페이지에서 다시 밟아본다.
    통과하면 앞선 동작들이 남긴 상태가 원인이고, 그대로 실패하면 기능 자체 문제다.
    앞선 동작이 없으면(= 이미 최소 단위) 가릴 게 없으니 아무것도 하지 않는다.
    """
    steps = case["steps"]
    failed = next((s["n"] for s in result["steps"]
                   if s["phase"] == "test" and s["status"] == "FAIL"), None)
    if not failed:
        return None
    idx = failed - 1                                   # 실패 스텝의 0-기반 위치

    first_goto = next((s for s in steps if s.get("action") == "goto"), None)
    last_act = next((i for i in range(idx - 1, -1, -1)
                     if steps[i].get("action") in INTERACTION_ACTIONS), None)
    if first_goto is None or last_act is None:
        return None                                    # 되짚을 선행 동작이 없다

    # 직전 동작 ~ 실패 스텝까지(사이의 wait 등 포함). 그 앞의 동작들은 일부러 뺀다.
    replay = [first_goto] + steps[last_act:idx + 1]
    try:
        for n, st in enumerate(replay, 1):
            runner.run(st, case["id"] + "_solo", n)
    except Exception:
        # 마지막 스텝에서 깨졌으면 단독으로도 실패, 앞에서 깨졌으면 판정 불가
        return "단독 재현: 실패 → 연속 사용과 무관한 결함" if st is replay[-1] else None
    return "단독 재현: 통과 → 이어서 조작할 때만 발생"


def run_case(runner, case, blocklist, dry_run):
    """케이스 하나 실행. 결과 dict 반환."""
    cid = case["id"]
    result = {
        "id": cid,
        "title": case["title"],
        "requirement": case.get("requirement", ""),
        "source": case.get("source", ""),
        # 실패했을 때 개발팀에 넘길 건지 기획팀이 먼저 볼 건지를 가르는 값
        "kind": case.get("kind", "defect"),
        # 웍스 등록용 로그에 쓰는 값. area 는 묶음 제목, fix 는 수정 지시 한 줄
        "area": case.get("area", ""),
        "fix": case.get("fix", ""),
        "writes": case.get("writes", False),
        "approved": case.get("approved", False),
        "status": "PASS",
        "reason": "",
        "steps": [],
        "shots": [],
        "duration": 0,
    }
    t0 = time.time()

    # -- 게이트 1: 쓰기 승인 --
    if case.get("writes"):
        if dry_run:
            result.update(status="SKIPPED", reason="dry-run 모드 - 쓰기 케이스 건너뜀")
            return result
        if not case.get("approved"):
            result.update(status="SKIPPED", reason="쓰기 미승인 - 케이스에 approved:true 필요")
            return result
        if not case.get("cleanup"):
            result.update(status="BLOCKED", reason="쓰기 케이스인데 cleanup 스텝이 없음 - 실행 거부")
            return result

    # -- 게이트 2: blocklist --
    for step in case["steps"] + case.get("cleanup", []):
        hit = hits_blocklist(step, blocklist)
        if hit:
            result.update(status="BLOCKED",
                          reason="blocklist 패턴 '{}' 매칭 - 실행 거부".format(hit))
            return result

    shots_before = len(runner.shots)

    def exec_steps(steps, phase):
        for i, step in enumerate(steps, 1):
            label = "{} {}".format(
                step.get("action"),
                step.get("selector", step.get("url", step.get("name", "")))
            )
            act = step.get("action")
            write = is_write_step(step)
            # 승인된 케이스여도 dry-run이면 쓰기 스텝은 건너뜀
            if write and dry_run:
                result["steps"].append({"phase": phase, "n": i, "label": label,
                                        "action": act,
                                        "status": "SKIP", "write": True, "msg": "dry-run"})
                continue
            try:
                runner.run(step, cid, i)
                result["steps"].append({"phase": phase, "n": i, "label": label,
                                        "action": act,
                                        "status": "OK", "write": write, "msg": ""})
            except Exception as e:
                msg = str(e).split("\n")[0][:300]
                result["steps"].append({"phase": phase, "n": i, "label": label,
                                        "action": act,
                                        "status": "FAIL", "write": write, "msg": msg})
                raise

    try:
        exec_steps(case["steps"], "test")
    except SessionExpired:
        raise                      # 개별 케이스 실패로 묻지 않고 전체 실행을 멈춘다
    except Exception as e:
        result.update(status="FAIL", reason=str(e).split("\n")[0][:300])
        try:
            runner._shoot("{}_FAIL".format(cid))
        except Exception:
            pass
        # chains 케이스는 "이어서 했기 때문인지" 를 자동으로 가린다
        if case.get("chains"):
            diag = diagnose_chain(runner, case, result)
            if diag:
                result["diagnosis"] = diag
                result["reason"] += "  [{}]".format(diag)

    # -- 정리(cleanup)는 본 테스트 실패와 무관하게 항상 시도 --
    if case.get("cleanup") and not dry_run and case.get("approved"):
        try:
            exec_steps(case["cleanup"], "cleanup")
        except Exception as e:
            result["reason"] += " / [!] cleanup 실패 - 생성 데이터 수동 확인 필요: {}".format(str(e)[:150])
            if result["status"] == "PASS":
                result["status"] = "FAIL"

    result["duration"] = round(time.time() - t0, 1)
    result["shots"] = runner.shots[shots_before:]
    return result


# -- 작업 폴더 초기화 -------------------------------------------

BAT_TEMPLATE = """@echo off
setlocal enabledelayedexpansion
set PYTHONIOENCODING=cp949:replace
cd /d "%~dp0"

echo.
echo ===============================================
echo   QA 자동화 실행
echo ===============================================
echo.

set RUNNER="{runner}"

set N=0
for %%F in (specs\\*.json) do (
    echo %%~nF | findstr /b "_" > nul
    if errorlevel 1 (
        set /a N+=1
        echo   !N!. %%~nF
    )
)

if %N%==0 (
    echo   실행할 스펙이 없습니다.
    echo   Claude에게 qa-test 스킬로 스펙을 만들어달라고 요청하세요.
    echo.
    pause
    exit /b 1
)

echo.
set "PICK="
set /p "PICK=실행할 번호 입력: "
if not defined PICK goto :badpick

rem Re-scan and match by index. Array vars (!SPEC[n]!) break when the
rem typed value carries stray characters, so compare as strings instead.
rem Keep comments ASCII - Korean bytes here can break the cmd parser.
set "TARGET="
set I=0
for %%F in (specs\\*.json) do (
    echo %%~nF | findstr /b "_" > nul
    if errorlevel 1 (
        set /a I+=1
        if "!I!"=="!PICK!" set "TARGET=%%F"
    )
)
if not defined TARGET goto :badpick
goto :picked

:badpick
echo.
echo   잘못된 번호입니다.
pause
exit /b 1

:picked
echo.
echo   1. 읽기만 실행  (안전 - 쓰기 케이스는 건너뜀)
echo   2. 전체 실행    (승인된 쓰기 케이스까지 실행)
echo.
set "MODE="
set /p "MODE=모드 번호 입력 [기본 1]: "
if not defined MODE set "MODE=1"

set "OPT=--dry-run"
if not "%MODE%"=="2" goto :run

set "OPT="
echo.
echo   [!] 전체 실행은 운영 데이터에 쓰기가 발생합니다.
set "OK="
set /p "OK=계속하려면 Y 입력: "
if /i not "%OK%"=="Y" (
    echo   취소되었습니다.
    pause
    exit /b 0
)

:run
echo.
py %RUNNER% "!TARGET!" %OPT%
set RESULT=%errorlevel%

echo.
for /f "delims=" %%D in ('dir /b /ad /o-d output 2^>nul') do (
    start chrome "%cd%\\output\\%%D\\리포트.html"
    goto :opened
)
:opened

echo.
pause
exit /b %RESULT%
"""


COOKIES_TEMPLATE = """# 크롬 F12 > Application > Cookies 에서 표 전체를 선택해 복사한 뒤
# 이 줄들 아래에 그대로 붙여넣으세요. JSON 으로 고칠 필요 없습니다.
# 컬럼 순서: Name / Value / Domain / Path / ... (탭 구분)
#
# 스펙에는  "cookies": "cookies.txt"  를 넣습니다.
# 로그인 창을 띄울 수 없는 환경에서만 쓰세요. 보통은  --login  으로 충분합니다.
# 세션을 브라우저 정보(UA)에 묶는 사이트는 복사한 쿠키가 안 먹을 수 있습니다.
# 그때는 스펙 meta 에  "user_agent": "<쿠키를 복사한 브라우저의 navigator.userAgent>"  를 넣습니다.
"""


LOGIN_BAT_TEMPLATE = """@echo off
setlocal
set PYTHONIOENCODING=cp949:replace
cd /d "%~dp0"

echo.
echo ===============================================
echo   로그인 세션 저장
echo ===============================================
echo.
echo   브라우저가 열리면 로그인만 하세요.
echo   로그인이 확인되면 저장하고 창이 저절로 닫힙니다.
echo   사이트가 여러 개면 주소를 띄어쓰기로 이어서 입력하세요.
echo.

set /p TARGET=로그인 주소 입력:
if "%TARGET%"=="" (
    echo   주소가 없습니다.
    pause
    exit /b 1
)

py "{runner}" --login %TARGET% --ws "%~dp0."

echo.
pause
"""


LOGIN_TIMEOUT = 15 * 60

# 화면이 로그인된 상태인지 본다. 페이지 안에서 읽기만 하므로 탭을 새로 띄우지 않는다.
# 로그인 안 된 신호: 보이는 비밀번호 칸, 또는 글자가 딱 "로그인"인 링크·버튼.
# 리다이렉트 중간의 빈 화면을 로그인으로 오판하지 않도록 본문 길이도 본다.
LOGIN_STATE_JS = """(done) => {
  if (document.readyState !== 'complete' || !document.body) return {ready: false};
  const vis = (el) => {
    const r = el.getBoundingClientRect(), s = getComputedStyle(el);
    return r.width > 0 && r.height > 0 && s.visibility !== 'hidden' && s.display !== 'none';
  };
  const word = /^(로그인|login|log in|sign in|signin)$/i;
  const text = document.body.innerText || '';
  return {
    ready: true,
    pw: [...document.querySelectorAll('input[type=password]')].some(vis),
    loginLink: [...document.querySelectorAll('a,button,input[type=submit],input[type=button]')]
      .some((el) => vis(el) && word.test((el.innerText || el.value || '').trim())),
    len: text.length,
    done: done ? text.includes(done) : null,
  };
}"""


def _site(host):
    """같은 사이트 판정용. admin.bluepharmkorea.co.kr 과 bluepharmkorea.co.kr 을 같은 곳으로 본다."""
    labels = (host or "").lower().split(".")
    n = 3 if len(labels) >= 3 and labels[-2] in ("co", "or", "go", "ne", "ac", "re", "pe") else 2
    return ".".join(labels[-n:])


def _host(url):
    from urllib.parse import urlparse
    try:
        return urlparse(url).hostname or ""
    except Exception:
        return ""


def _logged_in(page, done_text):
    try:
        st = page.evaluate(LOGIN_STATE_JS, done_text or "")
    except Exception:
        return False
    if not st.get("ready"):
        return False
    if done_text:
        return bool(st.get("done"))
    return not st["pw"] and not st["loginLink"] and st["len"] >= 200


def _snapshot(ctx):
    """storage_state 형식으로 세션을 모은다.

    ctx.storage_state() 는 열려 있지 않은 출처의 localStorage 를 읽으려고 숨은 탭을 띄웠다
    닫는다. 헤드풀 창에서 이걸 주기적으로 부르면 화면이 계속 깜빡이므로 쓰지 않는다.
    쿠키는 ctx.cookies(), localStorage 는 지금 열린 탭에서만 읽는다.
    """
    origins = {}
    for pg in ctx.pages:
        if pg.is_closed():
            continue
        try:
            o = pg.evaluate("""() => ({origin: location.origin,
                localStorage: Object.keys(localStorage).map(k => ({name: k, value: localStorage.getItem(k)}))})""")
        except Exception:
            continue
        if o["origin"].startswith("http"):
            origins[o["origin"]] = o
    return {"cookies": ctx.cookies(), "origins": list(origins.values())}


def _ua_path(auth):
    return auth.with_name(auth.stem + ".meta.json")


def do_login(urls, ws, done_text=None):
    """로그인 창을 띄우고, 로그인이 끝나면 세션을 저장한 뒤 창을 스스로 닫는다.

    사용자 동선은 "로그인만 한다" 로 끝난다. 창을 직접 닫을 필요도, 콘솔로 돌아갈 필요도 없다.
    주소를 여러 개 주면 하나가 로그인되는 대로 같은 탭을 다음 주소로 보낸다
    (예: 프론트 + 어드민). 기존 auth.json 이 있으면 불러와서 시작하므로, 이미 살아 있는
    사이트는 바로 통과하고 빠진 사이트만 로그인하면 된다.

    로그인 당시 브라우저의 UA 를 auth.meta.json 에 같이 남긴다. 세션을 UA 에 묶어두는
    사이트가 있어서(블루닥), 헤드리스 실행 때 UA 가 달라지면 쿠키가 살아 있어도 비로그인으로
    보인다. 러너는 이 파일이 있으면 같은 UA 로 실행한다.
    """
    ws.mkdir(parents=True, exist_ok=True)
    auth = ws / "auth.json"
    print("[*] 로그인 창을 엽니다: {}".format(" -> ".join(urls)))
    print("    로그인만 하세요. 로그인이 확인되면 저장하고 창이 저절로 닫힙니다.")

    done_sites, snap, ua = [], None, None
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        ctx_args = {"viewport": {"width": 1600, "height": 1000}}
        if auth.exists():
            ctx_args["storage_state"] = str(auth)
        ctx = browser.new_context(**ctx_args)
        page = ctx.new_page()
        page.goto(urls[0], wait_until="domcontentloaded", timeout=120000)
        ua = page.evaluate("navigator.userAgent")

        idx, streak, start = 0, 0, time.time()
        try:
            while time.time() - start < LOGIN_TIMEOUT:
                pages = [pg for pg in ctx.pages if not pg.is_closed()]
                if not pages:
                    break                                   # 사용자가 창을 닫음
                try:
                    snap = _snapshot(ctx)
                except Exception:
                    break
                site = _site(_host(urls[idx]))
                hit = next((pg for pg in pages
                            if _site(_host(pg.url)) == site and _logged_in(pg, done_text)), None)
                streak = streak + 1 if hit else 0
                if streak >= 2:                             # 두 번 연속이어야 확정 (중간 화면 오판 방지)
                    print("[*] 로그인 확인: {}".format(_host(urls[idx])))
                    done_sites.append(urls[idx])
                    idx, streak = idx + 1, 0
                    if idx >= len(urls):
                        snap = _snapshot(ctx)
                        break
                    hit.goto(urls[idx], wait_until="domcontentloaded", timeout=120000)
                    print("    다음 로그인: {}".format(urls[idx]))
                time.sleep(1.5)
        except KeyboardInterrupt:
            pass
        try:
            browser.close()
        except Exception:
            pass

    if not snap or not snap["cookies"]:
        print("[!] 세션이 저장되지 않았습니다. 다시 실행해 로그인하세요.")
        return 1
    auth.write_text(json.dumps(snap, ensure_ascii=False), encoding="utf-8")
    _ua_path(auth).write_text(json.dumps({"user_agent": ua}, ensure_ascii=False), encoding="utf-8")
    print("[*] 저장 완료: {} (쿠키 {}건)".format(auth, len(snap["cookies"])))
    missing = [u for u in urls if u not in done_sites]
    if missing:
        print("[!] 로그인 확인 전에 창이 닫혔거나 시간이 지났습니다: {}".format(", ".join(missing)))
        print("    이 주소는 로그인이 안 된 상태로 저장됐을 수 있습니다.")
        return 1
    print('    스펙의 meta 에 "auth": "auth.json" 을 넣으면 이 세션으로 실행됩니다.')
    return 0


def init_workspace(target):
    """사용자 프로젝트에 qa_auto/ 작업 폴더를 만든다. 스킬 폴더는 건드리지 않는다."""
    ws = target / "qa_auto"
    for sub in ("specs", "cases", "fixtures", "output"):
        (ws / sub).mkdir(parents=True, exist_ok=True)

    bat = ws / "run_qa.bat"
    # 한국어 윈도우 콘솔(cp949) 기준으로 배치 파일을 쓴다.
    # UTF-8 로 쓰면 cmd 가 배치 본문을 cp949 로 읽어 메뉴 한글이 깨지고,
    # BOM 을 붙이면 첫 줄(@echo off)까지 깨진다. 배치는 cp949, 파이썬 출력은
    # PYTHONIOENCODING 으로 맞추는 조합이 가장 안정적이다.
    bat.write_text(BAT_TEMPLATE.format(runner=SKILL_DIR / "qa_run.py"),
                   encoding="cp949", errors="replace")

    lbat = ws / "login_qa.bat"
    lbat.write_text(LOGIN_BAT_TEMPLATE.format(runner=SKILL_DIR / "qa_run.py"),
                    encoding="cp949", errors="replace")

    ck = ws / "cookies.txt"
    if not ck.exists():
        ck.write_text(COOKIES_TEMPLATE, encoding="utf-8")

    # 로그인 세션과 운영 화면 스크린샷이 실수로 커밋되는 사고를 막는다.
    gi = ws / ".gitignore"
    if not gi.exists():
        gi.write_text(
            "# 실제 로그인 세션 - 절대 커밋 금지\n"
            "auth.json\n"
            "auth.meta.json\n"
            "cookies.txt\n"
            "cookies.json\n"
            "\n"
            "# 운영 화면 스크린샷과 리포트\n"
            "output/\n"
            "\n"
            "# 케이스와 스펙은 공유 대상이므로 제외하지 않는다\n"
            "!specs/\n"
            "!cases/\n",
            encoding="utf-8")

    print("[*] 작업 폴더 생성: {}".format(ws))
    print("    specs/    실행 스펙 JSON")
    print("    cases/    테스트 케이스 MD")
    print("    fixtures/ 업로드 테스트용 파일")
    print("    output/   리포트 + 스크린샷")
    print("    run_qa.bat    더블클릭 실행")
    print("    login_qa.bat  로그인 세션 저장 (터미널 없이)")
    print("    cookies.txt  (선택) 개발자도구 쿠키 표 붙여넣기용")
    print("")
    print("[*] 로그인이 필요한 화면이면:")
    print('    py "{}" --login <URL> --ws "{}"'.format(SKILL_DIR / "qa_run.py", ws))
    return 0


# -- 메인 ------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec", nargs="?", help="spec.json 경로")
    ap.add_argument("--init", metavar="DIR",
                    help="해당 폴더에 작업 폴더(qa_auto/)와 실행 배치파일을 만든다")
    ap.add_argument("--login", metavar="URL", nargs="+",
                    help="로그인 창을 띄워 세션을 auth.json 으로 저장한다. 여러 주소면 차례로 로그인")
    ap.add_argument("--done-text", metavar="TEXT",
                    help="--login 에서 로그인 완료로 볼 문구 (예: 로그아웃). 없으면 화면으로 자동 판정")
    ap.add_argument("--ws", metavar="DIR", default="qa_auto",
                    help="--login 이 저장할 작업 폴더 (기본: ./qa_auto)")
    ap.add_argument("--dry-run", action="store_true", help="쓰기 스텝 전부 건너뜀")
    ap.add_argument("--headed", action="store_true", help="브라우저 표시")
    ap.add_argument("--case", help="특정 케이스 ID만 실행")
    args = ap.parse_args()

    if args.init:
        return init_workspace(Path(args.init).resolve())
    if args.login:
        return do_login(args.login, Path(args.ws).resolve(), args.done_text)
    if not args.spec:
        ap.error("spec 경로가 필요하다 (또는 --init DIR / --login URL)")

    spec_path = Path(args.spec).resolve()
    if not spec_path.exists():
        spec_path = (SKILL_DIR / args.spec).resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))

    workspace = resolve_workspace(spec_path)
    meta = spec["meta"]
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = workspace / "output" / "{}_{}".format(meta["project"], stamp)
    shot_dir = out_dir / "screenshots"
    shot_dir.mkdir(parents=True, exist_ok=True)

    cases = spec["cases"]
    if args.case:
        cases = [c for c in cases if c["id"] == args.case]
        if not cases:
            print("[!] 케이스 없음: {}".format(args.case))
            sys.exit(1)

    print("[*] 대상: {}".format(meta["base_url"].replace("{QA_ROOT}", SKILL_DIR.as_uri())))
    print("[*] 작업폴더: {}".format(workspace))
    print("[*] 케이스: {}건{}".format(
        len(cases), "  (DRY-RUN - 쓰기 없음)" if args.dry_run else ""))

    results = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not args.headed)
        ctx_args = {"viewport": meta.get("viewport", DEFAULT_VIEWPORT)}

        # auth: --login 으로 저장한 세션(쿠키 + localStorage). 쿠키 복붙보다 이쪽을 권한다.
        if meta.get("auth"):
            auth_path = None
            for base in (spec_path.parent, workspace):
                cand = (base / meta["auth"]).resolve()
                if cand.exists():
                    auth_path = cand
                    break
            if auth_path is None:
                print("[!] 인증 파일 없음: {} — `--login <URL>` 로 먼저 로그인하세요.".format(meta["auth"]))
                sys.exit(1)
            ctx_args["storage_state"] = str(auth_path)
            print("[*] 세션 사용: {}".format(auth_path.name))
            # 세션을 UA 에 묶는 사이트가 있어 로그인 당시 UA 로 맞춘다 (do_login 참고)
            ua_file = _ua_path(auth_path)
            if not meta.get("user_agent") and ua_file.exists():
                saved_ua = json.loads(ua_file.read_text(encoding="utf-8")).get("user_agent")
                if saved_ua:
                    ctx_args["user_agent"] = saved_ua

        if meta.get("user_agent"):
            ctx_args["user_agent"] = meta["user_agent"]

        ctx = browser.new_context(**ctx_args)
        if meta.get("cookies"):
            ctx.add_cookies(load_cookies(spec_path, meta["cookies"], workspace))
        page = ctx.new_page()
        runner = Runner(page, spec, spec_path, shot_dir, workspace)

        try:
            for case in cases:
                r = run_case(runner, case, spec.get("blocklist", []), args.dry_run)
                results.append(r)
                icon = {"PASS": "O", "FAIL": "X", "SKIPPED": "-", "BLOCKED": "!"}[r["status"]]
                print("  [{}] {} {}  {}".format(icon, r["id"], r["title"][:50], r["reason"][:60]))
        except SessionExpired as e:
            browser.close()
            print("\n[!] 실행 중단 - {}".format(e))
            print("    py qa_run.py --login <URL> --ws <작업폴더>  로 세션을 갱신하세요.")
            return 2

        browser.close()

    report_path = build_report(spec, results, runner.warnings, out_dir,
                               args.dry_run, str(spec_path))
    passed = sum(1 for r in results if r["status"] == "PASS")
    failed = sum(1 for r in results if r["status"] == "FAIL")
    print("\n[*] 결과: 통과 {} / 실패 {} / 전체 {}".format(passed, failed, len(results)))
    print("[*] 리포트: {}".format(report_path))
    n_dev = len(_dev_items(results))
    if n_dev:
        print("[*] 개발 전달용 {}건: {}".format(n_dev, out_dir / "개발전달.md"))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
